"""Build the human labeling workbook: one sheet per dev case, 41 CP rows with the
independent draft verdict pre-filled, plus a GOLD dropdown column for human verify."""
import os, re, json, csv
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT="/Users/siddhantparashar/projects/cyberai_task2"
CPS=json.load(open(f"{ROOT}/dev/cp_reference.json"))
DEV=list(csv.DictReader(open(f"{ROOT}/dev/dev_cases.csv")))

HDR=PatternFill("solid",fgColor="1F4E78"); HDRF=Font(color="FFFFFF",bold=True)
SUBF=Font(bold=True,color="1F4E78"); WRAP=Alignment(wrap_text=True,vertical="top")
FLAGFILL=PatternFill("solid",fgColor="FFF2CC"); GOLDFILL=PatternFill("solid",fgColor="E2EFDA")
thin=Side(style="thin",color="BBBBBB"); BORD=Border(thin,thin,thin,thin)

def scan_flags(evpath):
    t=open(evpath).read()
    flags=[]
    defs=len(re.findall(r'\[DEFICIENC',t,re.I))
    if defs: flags.append(f"{defs} explicit [DEFICIENCY] marker(s)")
    for lang in ["Mandarin","Vietnamese","Arabic","Chinese","Spanish"]:
        if re.search(r'\b'+lang+r'\b',t): flags.append(f"non-English mention: {lang}")
    if re.search(r'migration in progress|Jan 2025|>= 2 years|>=2 years',t): flags.append("possible retention/2-year issue (check dates)")
    for tr in range(1,10):
        if f"[Track {tr} was NOT provided" in t: flags.append(f"MISSING Track {tr}")
    return flags

def load_draft(case):
    p=f"{ROOT}/dev/drafts/{case}.json"
    if not os.path.exists(p): return {}
    try: d=json.load(open(p))
    except Exception: return {}
    return {v["cp"]:v for v in d.get("verdicts",[])}

wb=openpyxl.Workbook(); wb.remove(wb.active)
# README
rd=wb.create_sheet("README")
for i,line in enumerate([
 "FRECA Task 2 — DEV SET LABELING WORKBOOK",
 "",
 "One sheet per case (10 cases). Each has all 41 checking points.",
 "The 'DRAFT' columns are an INDEPENDENT model's draft (NOT the production pipeline) — treat as a starting hint only.",
 "YOUR JOB: fill the GOLD column (green) with the correct verdict: 1, 0, or N/A. Use the dropdown.",
 "Overwrite the draft whenever you disagree. Add reasoning in NOTES. Flag anything ambiguous.",
 "",
 "Verdict meaning:  1 = compliant · 0 = non-compliant/deficient · N/A = not applicable / cannot be assessed.",
 "Do NOT penalise cross-track commodity or establishment-name mismatches (that is template noise).",
 "The yellow 'FLAGS' box at the top of each sheet lists auto-detected things to look for — hints, not answers.",
 "",
 "Two people should label independently, then reconcile disagreements. Only reconciled cells are 'gold'.",
],1):
    c=rd.cell(row=i,column=1,value=line)
    if i==1: c.font=Font(bold=True,size=14)
rd.column_dimensions["A"].width=120

COLS=["CP","Element","Sub-element","Official checking-point text",
      "DRAFT verdict","DRAFT conf","DRAFT evidence","DRAFT reason",
      "GOLD verdict","Notes / reasoning"]
WIDTHS=[7,11,22,60,10,8,40,45,12,40]

for row in DEV:
    case=row["dev_id"]
    ws=wb.create_sheet(case[:31])
    draft=load_draft(case)
    flags=scan_flags(f"{ROOT}/dev/evidence/{case}.md")
    # header block
    ws.cell(row=1,column=1,value=f"CASE: {case}").font=Font(bold=True,size=13)
    ws.cell(row=2,column=1,value="FLAGS (hints, not answers): "+(" | ".join(flags) if flags else "none auto-detected"))
    ws.cell(row=2,column=1).fill=FLAGFILL
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=10)
    ws.cell(row=2,column=1).alignment=Alignment(wrap_text=True,vertical="top")
    ws.row_dimensions[2].height=42
    hr=4
    for j,(cname,w) in enumerate(zip(COLS,WIDTHS),1):
        c=ws.cell(row=hr,column=j,value=cname); c.fill=HDR; c.font=HDRF; c.alignment=WRAP; c.border=BORD
        ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes=f"A{hr+1}"
    cur_sub=None
    r=hr+1
    for cp in CPS:
        d=draft.get(cp["cp"],{})
        vals=[cp["cp"],cp["element"],cp["subelement"],cp["text"],
              d.get("verdict",""),d.get("confidence",""),d.get("evidence",""),d.get("reason",""),
              d.get("verdict",""),""]  # GOLD pre-seeded with draft to speed verify
        for j,v in enumerate(vals,1):
            c=ws.cell(row=r,column=j,value=v); c.alignment=WRAP; c.border=BORD
            if j==9: c.fill=GOLDFILL; c.font=Font(bold=True)
        r+=1
    # dropdown on GOLD column
    dv=DataValidation(type="list",formula1='"1,0,N/A"',allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"I{hr+1}:I{r-1}")

out=f"{ROOT}/dev/workbooks/FRECA_dev_labeling.xlsx"
wb.save(out)
print("wrote",out,"with",len(DEV),"case sheets")
