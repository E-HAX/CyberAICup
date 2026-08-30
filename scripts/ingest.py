"""Deterministic evidence ETL: one case folder -> evidence_context.md string.
Tolerant of missing/duplicated tracks. Preserves docx paragraph+table order and xlsx sheets."""
import os, re, sys
import docx
from docx.document import Document as _Doc
from docx.table import Table
from docx.text.paragraph import Paragraph
import openpyxl

def _iter_block_items(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag.endswith('}p'):
            yield Paragraph(child, parent)
        elif child.tag.endswith('}tbl'):
            yield Table(child, parent)

def docx_to_text(path):
    d = docx.Document(path)
    out = []
    for blk in _iter_block_items(d):
        if isinstance(blk, Paragraph):
            t = blk.text.strip()
            if t: out.append(t)
        else:
            for row in blk.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells): out.append(" | ".join(cells))
    return "\n".join(out)

def xlsx_to_text(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"### Sheet: {ws.title}")
        for r in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in r]
            if any(cells): out.append(" | ".join(cells))
    return "\n".join(out)

TRACK_NAMES = {1:"Establishment Registration Application Form",2:"HACCP Plan",
 3:"Pest Control Record",4:"Farm Management Plan",5:"Farm Site Plan",
 6:"Farm Hygiene & Sanitation Plan",7:"Bait Station Map",
 8:"Phytosanitary Security Procedure",9:"Traceability Records"}

def build_context(folder, only_substr=None):
    """only_substr: if set (e.g. 'Goldfields'), keep only files containing it (for the doubled folder)."""
    case_id = os.path.basename(folder.rstrip("/"))
    files = sorted(os.listdir(folder))
    if only_substr:
        files = [f for f in files if only_substr in f]
    by_track = {}
    for f in files:
        m = re.match(r'(\d+)_', f)
        if m: by_track.setdefault(int(m.group(1)), f)
    lines = [f"# CASE {case_id}",
             "Case identity anchors: the folder RE number above and the establishment name are authoritative.",
             "NOTE: evidence tracks are independently authored; per-track 'Registered Commodity', and any RE number / establishment name appearing INSIDE a track, are unreliable template fields and must NOT be treated as compliance evidence.",
             ""]
    for t in range(1, 10):
        lines.append(f"\n{'='*72}\n## Track {t} — {TRACK_NAMES[t]}\n{'='*72}")
        if t not in by_track:
            lines.append(f"[Track {t} was NOT provided for this case.]")
            continue
        p = os.path.join(folder, by_track[t])
        txt = docx_to_text(p) if p.endswith('.docx') else xlsx_to_text(p)
        lines.append(txt)
    return "\n".join(lines)

if __name__ == "__main__":
    folder = sys.argv[1]
    only = sys.argv[2] if len(sys.argv) > 2 else None
    print(build_context(folder, only))
